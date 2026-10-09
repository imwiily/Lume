import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from fonte.reader import Block
from fonte.editorial import analyze
from fonte.cli import main


def blocks(*texts):
    return [Block(i+1, text, 'Capítulo 1') for i, text in enumerate(texts)]


class EditorialTests(unittest.TestCase):
    def rules(self, *texts):
        return {f['rule'] for f in analyze(blocks(*texts))[0]}

    def test_duplicate(self):
        self.assertIn('frase_duplicada', self.rules('Ele começou a me explicar o mapa da região.', 'Ele começou a me explicar o mapa da região.'))

    def test_short_dialogue_not_duplicate(self):
        self.assertNotIn('frase_duplicada', self.rules('Sim, senhor.', 'Sim, senhor.'))

    def test_names(self):
        self.assertIn('variacao_nome', self.rules('Lívia entrou. Lívia sorriu.', 'Lívya voltou.'))

    def test_pronouns_and_interjections_are_not_names(self):
        self.assertNotIn('variacao_nome', self.rules('Você voltou. Você entrou. Vocês saíram.', 'Haaa! Haaa! Haaaa!', 'Hamm. Hamm. Humm. Humm.'))

    def test_different_evidence_changes_identity(self):
        # Mesmo alerta, mesma posição, evidência relacionada em outro texto: outra identidade.
        a = analyze(blocks('Lívia entrou. Lívia sorriu.', 'Lívya voltou.'))[0]
        b = analyze(blocks('Lívia entrou. Lívia sorriu.', 'Lívya voltou cedo.'))[0]
        self.assertEqual((a[0]['start'], a[0]['end']), (b[0]['start'], b[0]['end']))
        self.assertNotEqual(a[0]['id'], b[0]['id'])

    def test_retired_rules_are_accepted_and_silent(self):
        # Fase 7b: prazos de uma cena, repetição próxima e proximidade de referentes foram retirados.
        retiradas = {'duracao_suspensao', 'adiamento_amanha', 'palavra_proxima', 'referente_proximidade'}
        settings = {'rules': {r: True for r in retiradas}}
        textos = ('Não trabalha amanhã nem nos três dias seguintes.', 'O técnico foi afastado por três dias.',
                  'A entrega foi adiada para amanhã.', 'Ganhei mais dois dias.', 'Era melhor eu me preparar melhor.',
                  'A doutora se aproximou.', 'Eles estavam tão perto!')
        self.assertFalse(retiradas & {f['rule'] for f in analyze(blocks(*textos), settings=settings)[0]})

    def test_distinct_names(self):
        self.assertNotIn('variacao_nome', self.rules('Lívia entrou. Lívia sorriu.', 'Otávio voltou.'))

    def test_cut_and_reference(self):
        original = blocks('A médica se aproximou. Seus olhos azuis brilhavam intensamente sob a luz branca da sala.', 'Eles estavam tão perto!')
        revised = blocks('A médica se aproximou.', 'Eles estavam tão perto!')
        findings, _ = analyze(revised, original)
        scars = [f for f in findings if f['rule'] == 'pronome_apos_corte']
        self.assertEqual(len(scars), 1)
        self.assertEqual(scars[0]['related'][0]['document'], 'original')
        self.assertFalse(any(f['rule'] == 'pronome_apos_corte' for f in analyze(original, original)[0]))

    def test_unicode_and_stable_ids(self):
        bs = blocks('🌿 Cafe\u0301. Era melhor eu me preparar melhor.', 'Não trabalha amanhã nem nos três dias seguintes.', 'Foi afastado por três dias.')
        findings, _ = analyze(bs)
        self.assertEqual(findings, analyze(bs)[0])
        self.assertEqual(len(findings), len({f['id'] for f in findings}))
        for f in findings:
            self.assertTrue(0 <= f['start'] < f['end'] <= len(f['text']))
            for e in f['related'] + f['context']:
                self.assertTrue(0 <= e['start'] <= e['end'] <= len(e['text']))

    def test_cli_editorial_no_model_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            path = folder / "Livro d'água 🌿.docx"
            doc = Document()
            doc.add_paragraph('Ele começou a me explicar o mapa da região.')
            doc.add_paragraph('Ele começou a me explicar o mapa da região.')
            doc.save(path)
            before = path.read_bytes()
            output = folder / 'resultado'
            config = folder / 'busca.json'
            config.write_text(json.dumps({'rules': {
                'dialogo_contextual': False, 'referente_contextual': False, 'gerundismo': False,
                'memoria_narrativa': False, 'conflito_habilidade': False,
                'conflito_objeto': False, 'conflito_cronologia': False, 'coerencia_generica': False}}))
            args = ['revisar', str(path), '--modo', 'editorial', '--saida', str(output), '--config', str(config)]
            with patch('fonte.cli.load_model', side_effect=AssertionError('Editorial não deve carregar spaCy')):
                self.assertEqual(main(args), 0)
            report = json.loads((output / 'relatorio.json').read_text())
            self.assertEqual(report['metadata']['modo'], 'editorial')
            self.assertEqual(report['sha256'], hashlib.sha256(before).hexdigest())
            self.assertEqual(path.read_bytes(), before)
            self.assertIn('frase_duplicada', {f['rule'] for f in report['findings']})
            self.assertEqual(main(args), 2)
            self.assertEqual(path.read_bytes(), before)

    def test_cli_both_and_original(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            original, revised = folder/'original.docx', folder/'revisado.docx'
            doc = Document()
            doc.add_paragraph('A médica se aproximou. Seus olhos azuis brilhavam intensamente sob a luz branca da sala.')
            doc.add_paragraph('Eles estavam tão perto!')
            doc.save(original)
            doc.paragraphs[0].text = 'A médica se aproximou.'
            doc.save(revised)
            self.assertEqual(main(['revisar', str(revised), '--original', str(original), '--modo', 'ambas', '--saida', str(folder/'out')]), 0)
            report = json.loads((folder/'out/relatorio.json').read_text())
            self.assertEqual(report['metadata']['original_sha256'], hashlib.sha256(original.read_bytes()).hexdigest())
            self.assertTrue(any(f['rule'] == 'pronome_apos_corte' for f in report['findings']))

    def test_invalid_mode_combinations(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'test.docx'
            doc = Document(); doc.add_paragraph('Texto de teste.'); doc.save(path)
            self.assertEqual(main(['revisar', str(path), '--modo','editorial','--languagetool']), 2)
            self.assertEqual(main(['revisar', str(path), '--original', str(path)]), 2)

if __name__ == '__main__':
    unittest.main()


class ExpressiveRepetitionTests(unittest.TestCase):
    def test_onomatopoeia_reduplication_is_not_a_typo(self):
        from fonte.editorial.repetition import analyze as repetition

        def consecutive(text):
            return [f for f in repetition(blocks(text)) if f['rule'] == 'palavra_consecutiva']
        for text in ['— Au au, quieto aí, ninguém vai te machucar.', '— Blá blá, sempre a mesma história.']:
            with self.subTest(text=text):
                self.assertEqual(consecutive(text), [])
        self.assertTrue(consecutive('Ele saiu saiu de casa cedo.'))
