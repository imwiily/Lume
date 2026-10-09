from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fonte.reader import Block, read_docx, chapter_label
from fonte.settings import validate, DEFAULT
from fonte.segments import classify, masks
from fonte.editorial import analyze
from fonte.editorial.repetition import analyze as repetitions
from fonte.cli import main


def options(**kwargs):
    x=deepcopy(DEFAULT);x.update(kwargs);return x


class SearchSettingsTests(unittest.TestCase):
    def test_words_roman_digits(self):
        for label in ['Capítulo um','Capítulo dois','Capítulo três','Capítulo vinte e dois','Capítulo XI',' Capítulo 12 ','Capítulo primeiro','Capítulo um — A partida']:
            with self.subTest(label=label):self.assertTrue(chapter_label(label))
        for label in ['No capítulo um, saímos.', 'Capítulo estranho foi o de ontem.', 'Capítulo um foi escrito ontem.']:
            with self.subTest(label=label):self.assertFalse(chapter_label(label))

    def test_chapter_assignment_and_title_exclusion(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'book.docx';d=Document()
            for t in ['Capítulo um','O rapaz chegou.','Capítulo dois','Ele saiu.']:d.add_paragraph(t)
            d.save(path);bs,w=read_docx(path)
            self.assertEqual([b.chapter for b in bs],['Capítulo um']*2+['Capítulo dois']*2)
            self.assertTrue(bs[0].heading)
            self.assertFalse(any('Nenhum capítulo' in x for x in w))

    def test_custom_titles_and_auto_off(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'book.docx';d=Document()
            for t in ['CAPA','A floresta sem nome','Texto da história.']:d.add_paragraph(t)
            d.save(path);bs,_=read_docx(path, options(chapter_auto=False,chapter_titles=['A floresta sem nome']))
            self.assertFalse(bs[0].heading);self.assertTrue(bs[1].heading)
            self.assertEqual(bs[2].chapter,'A floresta sem nome')

    def test_outline_and_style(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'book.docx';d=Document();p=d.add_paragraph('A primeira noite')
            outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(outline)
            d.add_paragraph('Texto.');d.save(path)
            self.assertTrue(read_docx(path)[0][0].heading)
            self.assertFalse(read_docx(path,options(chapter_auto=False))[0][0].heading)
            self.assertTrue(read_docx(path,options(chapter_auto=False,chapter_styles=['Normal']))[0][0].heading)

    def test_no_chapter_warning(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'book.docx';d=Document();d.add_paragraph('Era uma noite fria.');d.save(path)
            self.assertTrue(any('Nenhum capítulo' in x for x in read_docx(path)[1]))

    def test_scope_dialogue_and_incise(self):
        b=Block(1,'— Calma! Calma! — pediu, tentando manter a calma.')
        found=repetitions([b],options(repetition_scopes=['narracao']))
        self.assertEqual(found,[])
        found=repetitions([b],options(repetition_scopes=['dialogo']))
        # Nunca usa o "calma" do inciso como segunda palavra.
        self.assertFalse(any(f['end']>b.text.index('pediu') for f in found))

    def test_no_word_crosses_incise(self):
        b=Block(1,'— Eu preciso da chave — disse, pegando a chave.')
        self.assertEqual(repetitions([b]),[])

    # Desde a Fase 7b, estas configurações são conferidas com a palavra dobrada: a repetição próxima,
    # usada antes como exemplo, foi retirada.
    def test_quotes_as_thoughts(self):
        b=Block(1,'“Ele saiu saiu de casa cedo.”')
        opts=options(quotes_role='pensamento',repetition_scopes=['narracao'])
        self.assertEqual(repetitions([b],opts),[])
        opts['repetition_scopes']=['pensamento']
        self.assertTrue(repetitions([b],opts))

    def test_italic_setting(self):
        text='Ele saiu saiu de casa cedo.';b=Block(1,text,italic=[(0,len(text))])
        self.assertEqual(repetitions([b],options(repetition_scopes=['narracao'])),[])
        self.assertTrue(repetitions([b],options(repetition_scopes=['narracao'],italic_thoughts=False)))

    def test_dashes_disabled(self):
        b=Block(1,'— Ele saiu saiu de casa cedo.')
        self.assertFalse(repetitions([b],options(repetition_scopes=['narracao'])))
        self.assertTrue(repetitions([b],options(repetition_scopes=['narracao'],dialogue_dashes=False)))

    def test_distance_and_boundary_are_still_accepted(self):
        # Só a repetição próxima (retirada na Fase 7b) usava a distância e o limite; as opções
        # continuam válidas nas configurações salvas, sem efeito sobre a palavra dobrada.
        b=Block(1,'Ele saiu saiu de casa cedo.')
        for opts in (options(word_distance=2), options(word_distance=40), options(repetition_boundary='frase')):
            self.assertEqual(len(repetitions([b],opts)),1)

    def test_duplicate_scope_and_exact_similar(self):
        bs=[Block(1,'Eu abri a pequena porta da casa antiga.'),Block(2,'Eu abri a pequena porta da casa vazia.')]
        self.assertFalse(any(f['rule']=='frase_duplicada' for f in repetitions(bs)))
        self.assertTrue(any(f['rule']=='frase_duplicada' for f in repetitions(bs,options(duplicate_similarity=.85))))
        self.assertFalse(any(f['rule']=='frase_duplicada' for f in repetitions(bs,options(duplicate_similarity=.85,duplicate_across_paragraphs=False))))

    def test_names_ignored_and_rule_off(self):
        bs=[Block(1,'Lívia chegou. Lívia sorriu. Lívya saiu.')]
        self.assertTrue(any(f['rule']=='variacao_nome' for f in analyze(bs)[0]))
        self.assertFalse(any(f['rule']=='variacao_nome' for f in analyze(bs,settings=options(ignored_names=['Lívia']))[0]))
        opts=options();opts['rules']={k:False for k in opts['rules']}
        self.assertEqual(analyze(bs,settings=opts)[0],[])

    def test_invalid_settings(self):
        for value in [{'word_distance':0},{'quotes_role':'auto'},{'rules':{'inexistente':True}},{'rules':{'tempo_verbal':'não'}},{'chapter_titles':'texto'},{'schema_version':2}]:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):validate(value)

    def test_offsets_preserved(self):
        b=Block(1,'🌿 “Eu vejo a rua.” Eu vejo a lua.')
        role=classify([b],options());masked=masks([b],role,['narracao'])[0]
        self.assertEqual(len(masked),len(b.text))
        self.assertTrue(masked.endswith('Eu vejo a lua.'))

    def test_cli_selective_without_model(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);path=root/'book.docx';d=Document()
            d.add_paragraph('Capítulo um');d.add_paragraph('Ele saiu saiu de casa cedo.');d.save(path)
            cfg=options();cfg['rules']={r:r=='palavra_consecutiva' for r in cfg['rules']}
            config=root/'busca.json';config.write_text(json.dumps(cfg))
            before=path.read_bytes()
            with patch('fonte.cli.load_model',side_effect=AssertionError('Não carregar modelo quando regras linguísticas desligadas')):
                self.assertEqual(main(['revisar',str(path),'--modo','ambas','--config',str(config),'--saida',str(root/'out')]),0)
            r=json.loads((root/'out/relatorio.json').read_text())
            self.assertEqual(len(r['findings']),1)
            self.assertEqual(r['findings'][0]['chapter'],'Capítulo um')
            self.assertEqual(r['metadata']['search_settings'],cfg)
            self.assertEqual(r['metadata']['chapters'][0]['title'],'Capítulo um')
            self.assertEqual(path.read_bytes(),before)

    def test_cli_all_off(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);path=root/'book.docx';d=Document();d.add_paragraph('Era melhor eu me preparar melhor.');d.save(path)
            cfg=options();cfg['rules']={r:False for r in cfg['rules']}
            config=root/'busca.json';config.write_text(json.dumps(cfg))
            self.assertEqual(main(['revisar',str(path),'--modo','ambas','--config',str(config),'--saida',str(root/'out')]),0)
            self.assertEqual(json.loads((root/'out/relatorio.json').read_text())['findings'],[])


class ModelScopeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import spacy
        cls.nlp=spacy.load('pt_core_news_sm',disable=['ner'])

    def test_tense_scopes(self):
        from fonte.search import linguistic
        bs=[Block(1,'Clara abriu a janela e observa a rua.'),Block(2,'— Clara abriu a janela e observa a rua.')]
        opts=options();opts['rules']={r:r=='tempo_verbal' for r in opts['rules']}
        f,_,_=linguistic(bs,self.nlp,'passado',opts)
        self.assertEqual({x['paragraph'] for x in f},{1})
        opts['tense_scopes']=['dialogo']
        f,_,_=linguistic(bs,self.nlp,'passado',opts)
        self.assertEqual({x['paragraph'] for x in f},{2})
        opts['tense_scopes']=[]
        self.assertEqual(linguistic(bs,self.nlp,'passado',opts)[0],[])

if __name__=='__main__':unittest.main()


class RetiredRulesTests(unittest.TestCase):
    def test_old_configs_with_narrative_memory_rules_still_load(self):
        from fonte.settings import RETIRED_RULES, RULES, validate
        ligado = validate({"rules": {**{r: True for r in RULES}, **{r: True for r in RETIRED_RULES}}})
        self.assertFalse(set(RETIRED_RULES) & set(ligado["rules"]))
        desligado = validate({"rules": {**{r: False for r in RULES}, **{r: False for r in RETIRED_RULES}}})
        self.assertFalse(any(desligado["rules"].values()))
        with self.assertRaises(ValueError):
            validate({"rules": {"regra_inexistente": True}})
