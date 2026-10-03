import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import spacy
from docx import Document

from fonte.analysis import analyze, narrative_masks
from fonte.lexicon import finite
from fonte.cli import main
from fonte.reader import Block, read_docx
from fonte.languagetool import check, utf16_index


class LinguisticTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.nlp = spacy.load("pt_core_news_sm", disable=["ner"])

    def scan(self, text, **kwargs):
        return analyze([Block(1, text)], self.nlp, kwargs.pop("tense", "passado"), **kwargs)[0]

    def test_present_in_past_narrative(self):
        f = self.scan("Davi andava pelas ruas vazias da cidade, quando alguém o intercepta por trás.")
        self.assertTrue(any(x["category"] == "Tempo verbal" and x["text"][x["start"]:x["end"]] == "intercepta" for x in f))

    def test_unseen_tense_example(self):
        f = self.scan("Clara abriu a janela e observa a rua.")
        self.assertTrue(any(x["category"] == "Tempo verbal" for x in f))

    def test_imperfect_counts_as_past(self):
        blocks = [Block(1,"Ele andava pela rua. Ana observava a casa. A menina caminhava sozinha.")]
        f, _, meta = analyze(blocks, self.nlp)
        self.assertEqual(meta["tempo"], "passado")
        self.assertEqual(f, [])

    def test_finite_auxiliaries(self):
        for text in ["Ele havia chegado cedo.", "Ela tinha visto a ave na árvore."]:
            self.assertEqual(self.scan(text), [])

    def test_dialogue_excluded(self):
        for text in ['“Eu não sei o que está acontecendo”, disse Davi.', '"Eu estou bem", respondeu Ana.', '— Eu estou bem — disse Davi. — Não se preocupe.']:
            self.assertEqual(self.scan(text), [])

    def test_multiline_quote(self):
        blocks = [Block(1, '“Eu estou bem.'), Block(2, 'Eu vejo a saída”, disse Davi.')]
        f, warnings, _ = analyze(blocks, self.nlp, "passado")
        self.assertEqual(f, [])
        self.assertEqual(warnings, [])

    def test_unclosed_quote_warns(self):
        _, _, warnings = narrative_masks([Block(1, '“Eu vejo tudo.')])
        self.assertTrue(warnings)

    def test_broken_quote_does_not_invert_later_dialogue(self):
        blocks = [Block(1, '\"Hm? O guarda franziu a testa.'),
                  Block(2, '\"Eu estou bem\"'), Block(3, 'Ele andava pela rua.')]
        f, warnings, _ = analyze(blocks, self.nlp, "passado")
        self.assertTrue(warnings)
        self.assertEqual(f, [])

    def test_nominal_fragment(self):
        f = self.scan("Um pedaço de papel seu rasgado.")
        self.assertEqual(f[0]["category"], "Estrutura da frase")

    def test_unseen_fragment(self):
        f = self.scan("Uma velha cadeira de madeira no canto da sala.")
        self.assertTrue(any(x["category"] == "Estrutura da frase" for x in f))

    def test_short_fragments_preserved(self):
        self.assertEqual(self.scan("Buzinas. Portas batendo. O vizinho reclamando."), [])

    def test_dialogue_action_comma(self):
        f = self.scan('“Inacreditável”, as mãos do guarda voltavam à posição normal.')
        self.assertTrue(any(x["category"] == "Pontuação de diálogo" for x in f))

    def test_dialogue_valid_action_separated(self):
        self.assertEqual(self.scan('“Inacreditável.” As mãos do guarda voltavam à posição normal.'), [])

    def test_italic_thought(self):
        text = "Eu vejo tudo agora."
        f, _, _ = analyze([Block(1,text,italic=[(0,len(text))])], self.nlp, "passado")
        self.assertEqual(f, [])

    def test_repeat_narrative_only(self):
        f = self.scan('Ele pegou a a bolsa. “Eu eu não sei”, disse Davi.')
        self.assertEqual(sum(x["category"] == "Palavra repetida" for x in f), 1)

    def test_chapter_and_offsets(self):
        text = "Ana caminhava pela praça quando encontra João."
        f, _, _ = analyze([Block(9,text,"Capítulo 2")], self.nlp, "passado")
        self.assertEqual(f[0]["paragraph"], 9)
        self.assertEqual(f[0]["chapter"], "Capítulo 2")
        self.assertEqual(text[f[0]["start"]:f[0]["end"]], "encontra")

    def test_reader_tables_italics_and_input_unchanged(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td)/"Teste.docx"
            d=Document();d.add_heading("Capítulo 1",1);d.add_paragraph('')
            p=d.add_paragraph();p.add_run("Pensamento.").italic=True
            d.add_table(1,1).cell(0,0).text="Texto na tabela."
            d.save(path)
            digest=hashlib.sha256(path.read_bytes()).hexdigest()
            blocks, _ = read_docx(path)
            self.assertEqual(blocks[-1].text,"Texto na tabela.")
            self.assertEqual(blocks[1].italic,[(0,11)])
            self.assertEqual(blocks[1].number,3)
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),digest)

    def test_cli_integration_preserves_docx(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/"Manuscrito com espaços.docx";out=Path(td)/"resultado"
            d=Document();d.add_paragraph("Davi andava pela rua quando alguém o intercepta por trás.");d.save(path)
            original=path.read_bytes()
            with patch("fonte.cli.load_model",return_value=self.nlp):
                self.assertEqual(main(["revisar",str(path),"--saida",str(out),"--tempo","passado"]),0)
                self.assertEqual(main(["revisar",str(path),"--saida",str(out)]),2)
            self.assertEqual(path.read_bytes(),original)
            data=json.loads((out/"relatorio.json").read_text())
            self.assertTrue(data["findings"])
            self.assertEqual(sorted(p.name for p in out.iterdir()),["relatorio.json"])

    def test_pages_rejected(self):
        with self.assertRaisesRegex(ValueError, "Pages"):
            read_docx(Path("x.pages"))


class InfrastructureTests(unittest.TestCase):
    def test_utf16_emoji(self):
        self.assertEqual(utf16_index("😀 palavra",3),2)

    def test_languagetool_contract(self):
        # Simula o contrato HTTP; não valida a qualidade das regras do LT.
        class Response:
            def __enter__(self):return self
            def __exit__(self,*args):pass
            def read(self):return json.dumps({"matches":[{"offset":3,"length":7,"message":"Confira a grafia","rule":{"id":"MORFOLOGIK_RULE_PT_BR","issueType":"misspelling"}}]}).encode()
        with patch("fonte.languagetool.build_opener") as builder:
            builder.return_value.open.return_value=Response()
            results,_=check([Block(1,"😀 dividos")])
            self.assertEqual(results[0]["text"][results[0]["start"]:results[0]["end"]],"dividos")
            request=builder.return_value.open.call_args.args[0]
            self.assertEqual(request.full_url,"http://127.0.0.1:8081/v2/check")


if __name__ == "__main__":
    unittest.main()


class FragmentVerbParaTests(unittest.TestCase):
    """“para” seguido de “no”, “de”, “em”… só pode ser o verbo parar."""
    @classmethod
    def setUpClass(cls):
        cls.nlp = spacy.load("pt_core_news_sm", disable=["ner"])

    scan = LinguisticTests.scan

    def fragments(self, text, tense="passado"):
        return [x for x in self.scan(text, tense=tense) if x["category"] == "Estrutura da frase"]

    def test_verb_para_is_finite(self):
        for text in ["Ela tenta de novo. Mas a velha carroça de madeira para no meio da subida.",
                     "Ele corre muito. Mas o velho relógio da praça para de repente outra vez.",
                     "Tudo segue igual. O ônibus lotado da manhã para em frente à escola antiga."]:
            with self.subTest(text=text):
                self.assertEqual(self.fragments(text, "presente"), [])

    def test_preposition_para_keeps_nominal_fragment(self):
        self.assertTrue(self.fragments("Uma longa viagem de trem para o norte, sem destino algum."))


class FragmentEllipsisTests(unittest.TestCase):
    """Complemento solto que retoma o verbo da frase anterior (“[pensei] naquela…”)."""
    @classmethod
    def setUpClass(cls):
        cls.nlp = spacy.load("pt_core_news_sm", disable=["ner"])

    scan = LinguisticTests.scan
    fragments = FragmentVerbParaTests.fragments

    def test_parallel_complement_recovers_previous_verb(self):
        for text in ["Pensei no jardim molhado. Naquela mulher de chapéu azul. Ela sorria.",
                     "Lembrei da casa antiga. Daquele homem de casaco cinza.",
                     "Confiei no velho capitão. No mapa amarelado do avô dele.",
                     "Ele falava com a irmã mais nova. Com o cachorro do vizinho também.",
                     "Corremos pela estrada de terra. Pelos campos molhados de chuva fina."]:
            with self.subTest(text=text):
                self.assertEqual(self.fragments(text), [])

    def test_without_recoverable_verb_still_candidate(self):
        for text in ["Do outro lado, um grupo de turistas.",
                     # Preposição diferente: “pensei daquele…” não é retomada.
                     "Pensei no jardim molhado. Daquele homem de casaco cinza escuro.",
                     # (“Uma tarde inteira de chuva. Na janela…” saiu daqui: sequência de fragmentos
                     # descritivos deixou de gerar alerta; ver FragmentClassificationTests.)
                     # Adjunto anteposto e núcleo nominal: não é complemento solto.
                     "Pensei no jardim molhado. No canto, um velho banco de pedra."]:
            with self.subTest(text=text):
                self.assertTrue(self.fragments(text))

    def test_reason_names_segment_without_technical_disclaimer(self):
        f, = self.fragments("Uma velha cadeira de madeira no canto da sala.")
        self.assertIn("‘Uma velha cadeira de madeira no canto da sala’", f["reason"])
        self.assertNotIn("analisador", f["reason"])


class FragmentSuspensionTests(unittest.TestCase):
    """Reticências suspendem o pensamento; segue uma constatação nominal curta."""
    @classmethod
    def setUpClass(cls):
        cls.nlp = spacy.load("pt_core_news_sm", disable=["ner"])

    scan = LinguisticTests.scan
    fragments = FragmentVerbParaTests.fragments

    def test_short_nominal_after_suspension_is_not_reported(self):
        for text in ["Eu queria explicar tudo. Mas, depois daquela prova difícil… nota final zero.",
                     "E aquela expressão no rosto dele… puro medo.",
                     "Depois do exame de sangue da manhã... nenhuma resposta.",
                     "Mas depois daquela noite tão longa… silêncio absoluto."]:
            with self.subTest(text=text):
                self.assertEqual(self.fragments(text), [])

    def test_long_continuation_or_no_suspension_still_candidate(self):
        for text in ["Depois daquele dia… uma longa fila de carros parados na avenida principal da cidade.",
                     "Uma velha cadeira de madeira no canto da sala."]:
            with self.subTest(text=text):
                self.assertTrue(self.fragments(text))


class FragmentClassificationTests(unittest.TestCase):
    """Fragmento sem verbo finito: só a oração incompleta gera alerta normal; o incerto, de baixa confiança."""
    @classmethod
    def setUpClass(cls):
        cls.nlp = spacy.load("pt_core_news_sm", disable=["ner"])

    scan = LinguisticTests.scan
    fragments = FragmentVerbParaTests.fragments

    def test_elliptical_predication_recovers_previous_referent(self):
        # O adjetivo concorda com um referente da frase anterior: “[as chaves eram] enferrujadas…”.
        for text in ["Encontrei duas chaves antigas na gaveta. Enferrujadas e tortas por causa da umidade.",
                     "Vi os barcos no cais. Minúsculos diante do enorme navio do porto.",
                     "Vi o animal no fundo do poço. Enorme, escuro e completamente imóvel.",
                     "Encontramos três garrafas na areia molhada. Verdes. Minúsculas diante da força das ondas."]:
            with self.subTest(text=text):
                self.assertEqual(self.fragments(text), [])

    def test_adverbial_fragment_continues_previous_action(self):
        for text in ["O ponteiro começou a girar. Primeiro bem devagar e, em seguida, mais e mais depressa.",
                     "A luz vinha na nossa direção. Aos poucos, cada vez mais perto da margem.",
                     "Ela correu sem olhar para trás. Diretamente para o portão da fazenda."]:
            with self.subTest(text=text):
                self.assertEqual(self.fragments(text), [])

    def test_descriptive_sequence_and_short_fragments(self):
        for text in ["O quarto estava vazio. Silêncio por toda a casa antiga.",
                     "Uma tarde inteira de chuva. Na janela da sala de jantar.",
                     "Vi a criatura. Enorme. Escura. Imóvel.",
                     "Ele abriu a porta. Escuridão absoluta.",
                     "Aos poucos. Cada vez mais perto."]:
            with self.subTest(text=text):
                self.assertEqual(self.fragments(text), [])

    def test_incomplete_clause_keeps_normal_alert(self):
        # Subordinante ou relativo sem verbo: falta a oração principal.
        for text, marked in [("Enquanto todos na sala, depois do jantar longo.", "Enquanto todos na sala, depois do jantar longo"),
                             ("A menina que, sentada perto da janela aberta.", "A menina que, sentada perto da janela aberta")]:
            with self.subTest(text=text):
                f, = self.fragments(text)
                self.assertEqual(f["text"][f["start"]:f["end"]], marked)
                self.assertEqual(f.get("confidence", "média"), "média")
                self.assertIn("incompleta", f["reason"])

    def test_uncertain_nominal_fragment_is_low_confidence(self):
        # Frase nominal isolada, sem apoio no contexto: continua visível, com confiança baixa.
        for text in ["Uma velha cadeira de madeira no canto da sala.",
                     # Sem concordância com a frase anterior, o adjetivo não recupera referente.
                     "O carro parou na esquina. Minúsculas diante do tamanho da praça."]:
            with self.subTest(text=text):
                f, = self.fragments(text)
                self.assertEqual((f["confidence"], f["confidence_score"]), ("baixa", .4))

    def test_classes_are_counted_in_metadata(self):
        text = "Uma velha cadeira de madeira no canto da sala. Vi os barcos no cais. Minúsculos diante do enorme navio do porto."
        counts = analyze([Block(1, text)], self.nlp, "passado")[2]["fragmentos_sem_verbo"]
        self.assertEqual(counts, {"likely_elliptical_predication": 1, "uncertain": 1})


class FiniteVerbTests(unittest.TestCase):
    """Formas finitas de verbos de ligação, auxiliares e lexicais, em vários tempos."""
    @classmethod
    def setUpClass(cls):
        cls.nlp = spacy.load("pt_core_news_sm", disable=["ner"])

    scan = LinguisticTests.scan
    fragments = FragmentVerbParaTests.fragments

    def has_finite(self, text):
        return any(finite(t) for t in self.nlp(text))

    def test_copula_tagged_as_verb_or_adverb_is_finite(self):
        # O modelo liga a cópula ao predicativo (dep=cop) e a etiqueta como VERB ou ADV; “era”
        # também é substantivo no léxico. A relação sintática decide, não só a etiqueta.
        for text in ["Não era como uma ponte qualquer.", "Era como uma porta antiga.", "Era tarde demais."]:
            with self.subTest(text=text):
                self.assertTrue(self.has_finite(text))
        self.assertEqual(self.fragments("Não era como uma ponte qualquer."), [])

    def test_paradigms_are_finite(self):
        frases = (["Ele {} muito alto.".format(f) for f in "é era foi será seria".split()]
                  + ["Talvez ele seja muito alto.", "Se ele fosse muito alto, entraria."]
                  + ["Ela {} cansada.".format(f) for f in "está estava esteve estará estaria".split()]
                  + ["Talvez ela esteja cansada.", "Se ela estivesse cansada, dormiria."]
                  + ["Ele {} medo do escuro.".format(f) for f in "tem tinha teve terá teria".split()]
                  + ["{} alguém na porta.".format(f) for f in "Há Havia Houve Haverá Haveria".split()]
                  + ["Ele {} pela praia.".format(f) for f in "corre correu corria correrá correria".split()]
                  + ["Eles foram embora.", "Tenho medo.", "Ele tinha saído cedo.", "Pode acontecer.",
                     "Deveria funcionar.", "Ficou parado.", "Parece estranho.", "Continuava escuro lá fora."])
        for text in frases:
            with self.subTest(text=text):
                self.assertTrue(self.has_finite(text))

    def test_nonfinite_forms_alone_are_not_finite(self):
        for text in ["correr pela floresta", "ter terminado o trabalho", "sendo observado", "feito por ele",
                     "ao chegar em casa"]:
            with self.subTest(text=text):
                self.assertFalse(self.has_finite(text))
