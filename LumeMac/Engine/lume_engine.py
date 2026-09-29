"""Entrada única do executável autossuficiente usado pelo Lume."""
import json
from pathlib import Path
import sys


def model():
    import pt_core_news_sm
    return pt_core_news_sm.load(disable=['ner'])


def health():
    from fonte import __version__
    from fonte.analysis import analyze
    from fonte.reader import read_docx
    from fonte.report import render
    from docx import Document
    import tempfile
    with tempfile.TemporaryDirectory() as temporary:
        path = Path(temporary) / 'teste.docx'
        document = Document(); document.add_paragraph('Capítulo um'); document.add_paragraph('Clara abriu a janela e observa a rua.'); document.save(path)
        blocks, _ = read_docx(path)
        assert blocks[0].heading and blocks[1].chapter == 'Capítulo um', 'Identificação de capítulo falhou'
        findings, _, _ = analyze(blocks, model(), 'passado', True)
        assert any(f['text'][f['start']:f['end']] == 'observa' for f in findings), 'Análise de teste falhou'
        assert '__REPORT_DATA__' not in render({'findings': findings}), 'Template HTML ausente'
        from fonte.editorial import analyze as editorial
        from fonte.reader import Block
        editorial_findings, _ = editorial([Block(1, 'Não virá amanhã nem nos três dias seguintes. Está suspenso por três dias.')])
        assert any(f['rule'] == 'duracao_suspensao' for f in editorial_findings), 'Análise editorial ausente'
        from fonte.settings import validate as search_settings
        settings = search_settings({'rules': {'duracao_suspensao': False}})
        filtered, _ = editorial([Block(1, 'Não virá amanhã nem nos três dias seguintes. Está suspenso por três dias.')], settings=settings)
        assert not any(f['rule'] == 'duracao_suspensao' for f in filtered), 'Filtro editorial falhou'
        from fonte.pipeline import run
        from fonte.settings import NEW_RULES, validate
        options = validate({})
        options['rules'] = {name: name in NEW_RULES for name in options['rules']}
        modular, _, meta = run([Block(1, '🌿 Nada além de disso.'),
                               Block(2, 'O artesão fabricaria um vaso que receberá pinturas.')],
                              model, settings=options, tense='passado')
        assert modular[0]['excerpt'] == 'além de disso' and modular[0]['module'] == 'linguistic'
        relation = next(f for f in modular if f.get('relation') == 'conditional_future')
        anchor = relation['related'][0]
        assert relation['suggestion'] == 'receberia' and anchor['text'][anchor['start']:anchor['end']] == 'fabricaria', 'Relação temporal ausente'
        semantic, _, memory = run([Block(1, 'Íris só usa fogo.'), Block(2, 'Íris criou gelo.')], model, mode='editorial')
        assert any(f.get('category_code') == 'ability_conflict' for f in semantic), 'Memória narrativa ausente'
        assert memory['fact_bank']['facts'], 'Banco de fatos ausente'
        quality, _, upgraded = run([Block(1, 'Íris dominava apenas magia de fogo.'),
                                   Block(2, 'Ela concentrou magia e criou gelo.'),
                                   Block(3, '— Não vou — disse ela.'),
                                   Block(4, 'O dispositivo de Íris foi destruído.'),
                                   Block(5, 'O dispositivo de Tomás estava intacto.')], model, mode='editorial')
        facts = upgraded['fact_bank']['facts']
        fire = next(f for f in facts if f['relation'] == 'ability' and f['scope'] == 'exclusive')
        ice = next(f for f in facts if f['relation'] == 'demonstrated_ability')
        assert fire['value'] == 'fogo' and ice['value'] == 'gelo' and fire['subject'] == ice['subject'], 'Fatos fogo/gelo inválidos'
        assert any(e['type'] == 'USE_ABILITY' and e['id'] == ice['event_id'] for e in upgraded['fact_bank']['events']), 'Evento de habilidade ausente'
        assert upgraded['scenes'][0]['dialogue_turns'][0]['speaker'] == fire['subject'], 'Falante pronominal ausente'
        assert len(upgraded['fact_bank']['objects']) == 2, 'Objetos homônimos indevidamente unidos'
        assert not any(f.get('category_code') == 'object_state_conflict' for f in quality), 'Conflito entre identidades distintas'

        generic, _, _ = run([Block(1, 'João tinha 32 anos.'),
                              Block(2, 'João completou 29 anos.')], model, mode='editorial')
        assert any(f.get('category_code') == 'age_conflict' for f in generic), 'Fatos genéricos ausentes'
        controls, _, _ = run([Block(1, 'A porta estava fechada.'), Block(2, 'Ana abriu a porta.'),
                               Block(3, 'A porta estava aberta.')], model, mode='editorial')
        assert not any(f.get('rule') == 'coerencia_generica' for f in controls), 'Transição legítima gerou conflito'

        from fonte.semantic import key
        _, _, roles = run([Block(1, 'A chave foi entregue a Pedro por Ana.'),
                           Block(2, 'Ana tinha 32 cartas.')], model, mode='editorial')
        role_facts = roles['fact_bank']['facts']
        holder = next(f for f in role_facts if f['relation'] == 'holder')
        transfer = next(e for e in roles['fact_bank']['events'] if e['id'] == holder['event_id'])
        assert transfer['subject'] == key('char', 'Ana') and holder['value'] == key('char', 'Pedro'), 'Papéis da transferência incorretos'
        assert not any(f['relation'] == 'age' for f in role_facts), 'Quantidade confundida com idade'

        _, _, real = run([Block(1, 'Maria entrou.'),
                          Block(2, '“Estou bem”, disse Maria, acenando com a cabeça.'),
                          Block(3, 'Maria tinha trinta e quatro anos.')], model, mode='editorial')
        assert real['scenes'][0]['dialogue_turns'][0]['speaker'] == key('char','Maria'), 'Falante entre aspas ausente'
        assert not any(f['relation']=='tells' for f in real['fact_bank']['facts']), 'Gesto confundido com conteúdo de fala'
        assert any(f['relation']=='age' and f['value']==34 for f in real['fact_bank']['facts']), 'Numeral composto não reconhecido'
        _, _, first = run([Block(1, 'Cheguei em casa.'), Block(2, 'Abri a porta.')], model, mode='editorial')
        assert first['scenes'][0]['narrator'] != 'unknown', 'Narrador em primeira pessoa ausente'
        assert any(f['relation']=='door_state' and f['value']=='open' for f in first['fact_bank']['facts']), 'Estado da porta ausente'

        clitic, _, _ = run([Block(1, '— Não vou — virou-se Helena.')], model, mode='editorial')
        assert any(f.get('category_code') == 'narrative_action_after_speech' and f['excerpt'] == 'virou-se' for f in clitic), 'Intervalo de ação pronominal inválido'
        assert meta['stages'][-1]['state'] == 'not_implemented', 'Auditoria indevidamente anunciada'
        grammar, _, _ = run([Block(1, 'Ela entregou o livro a professora.')], model, tense='passado')
        assert any(f.get('rule') == 'crase' for f in grammar), 'Regras gramaticais ausentes'
        from fonte.coerencia_ia import estimar
        estimativa = estimar([Block(1, 'Capítulo 1', 'Capítulo 1', heading=True), Block(2, 'Lia tinha olhos verdes.', 'Capítulo 1')],
                             Path(temporary) / 'coerencia')
        assert estimativa['a_enviar'] == 1 and estimativa['custo_estimado_usd'] > 0, 'Coerência com IA ausente'
    from fonte.languagetool import available
    return {'api_version': 1, 'report_schema': 1, 'decision_schema': 1,
            'engine_version': __version__, 'healthy': True, 'grammar_checker': available(), 'coherence_ai': True}


def main():
    args = sys.argv[1:]
    if args == ['--lume-probe']:
        print(json.dumps(health())); return 0
    if args and args[0] in ('--lume-install', '--lume-rollback', '--lume-reset'):
        import argparse
        from engine_packages import install, switch
        parser = argparse.ArgumentParser()
        parser.add_argument('--support', required=True)
        parser.add_argument('--package')
        parsed = parser.parse_args(args[1:])
        if args[0] == '--lume-install':
            if not parsed.package: raise ValueError('Selecione o pacote de atualização.')
            result = install(parsed.package, parsed.support)
        else:
            root = Path(sys.executable).resolve().parent.parent
            result = switch(parsed.support, root, reset=args[0] == '--lume-reset')
        print(json.dumps({'engine_version': result['engine_version'], 'status': 'ok'})); return 0
    from fonte import cli
    cli.load_model = model
    return cli.main(args)


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except Exception as error:
        print('Não foi possível concluir: ' + str(error), file=sys.stderr)
        raise SystemExit(2)
